"""Adaptador de entrega: DOCX (python-docx).

Layout limpo com seções por setor. write() gera o arquivo e devolve o caminho.
Testável sozinho: basta chamar write() com um Briefing.
"""
from __future__ import annotations

import logging
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

from impact.models import Briefing

log = logging.getLogger("delivery.docx")

_SEV_COR = {
    1: RGBColor(0x2E, 0x7D, 0x32), 2: RGBColor(0x2E, 0x7D, 0x32),
    3: RGBColor(0xF9, 0xA8, 0x25), 4: RGBColor(0xEF, 0x6C, 0x00),
    5: RGBColor(0xC6, 0x28, 0x28),
}


def write(briefing: Briefing, path: str | Path = "out/briefing.docx") -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading(briefing.titulo, level=0)
    p = doc.add_paragraph()
    p.add_run(
        f"Gerado em {briefing.gerado_em} · {len(briefing.eventos)} eventos · "
        f"modo {briefing.modo} · {briefing.total_ingerido} eventos ingeridos"
    ).italic = True

    for setor, eventos in sorted(briefing.por_setor.items()):
        editoria = getattr(eventos[0], "editoria", "") or setor
        doc.add_heading(editoria, level=1)
        for ev in eventos:
            h = doc.add_heading(level=2)
            run = h.add_run(f"[sev {ev.severidade}/5] {ev.titulo}")
            run.font.color.rgb = _SEV_COR.get(ev.severidade, RGBColor(0, 0, 0))

            meta = doc.add_paragraph()
            meta.add_run(
                f"{ev.geografia} · relevância {ev.relevancia}/10 · fonte: {ev.fonte}"
            ).italic = True

            if getattr(ev, "resumo", ""):
                doc.add_paragraph(ev.resumo)
            if getattr(ev, "projecao", ""):
                pj = doc.add_paragraph()
                pj.add_run("Cenário provável: ").bold = True
                pj.add_run(ev.projecao)

            for label, valor in (
                ("Impacto no custo", ev.impacto_custo),
                ("Impacto na cadeia", ev.impacto_cadeia),
                ("Exposição cyber", ev.exposicao_cyber),
                ("Ação recomendada", ev.acao_recomendada),
            ):
                par = doc.add_paragraph(style="List Bullet")
                par.add_run(f"{label}: ").bold = True
                par.add_run(valor)

            orig = doc.add_paragraph()
            idi = (getattr(ev, "idioma_original", "") or "").upper()
            orig.add_run(
                f"Fonte original ({idi}): {getattr(ev, 'titulo_original', '') or ev.titulo}"
            ).italic = True
            orig.runs[0].font.size = Pt(8)
            if ev.url:
                link = doc.add_paragraph()
                link.add_run(ev.url).font.size = Pt(8)

    doc.save(str(path))
    log.info("DOCX gerado em %s", path)
    return path
